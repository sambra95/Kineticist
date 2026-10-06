"""Reading plate-reader kinetic exports (.xlsx, one plate per file)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime, time, timedelta

import numpy as np
import openpyxl

ROWS = list("ABCDEFGH")
COLS = list(range(1, 13))
WELLS = [f"{r}{c}" for r in ROWS for c in COLS]  # row-major A1..A12, B1..B12, ...

@dataclass
class Plate:
    """One plate's kinetic read. ``a`` is (n_timepoints, 96) in WELLS order."""

    file: str
    t: np.ndarray                     # seconds from the start of the read
    temp: np.ndarray                  # incubator temperature per read, deg C
    a: np.ndarray                     # absorbance; NaN where saturated or missing
    sat: np.ndarray                   # True where the reader reported OVRFLW
    dropped: int                      # trailing partial sweeps removed
    started: datetime | None          # date and time the read began

    @property
    def clock(self) -> str:
        return self.started.strftime("%H:%M:%S") if self.started else ""


def _seconds(v) -> float:
    """Kinetic timestamps come back as datetime.time, or timedelta past 24 h."""
    if v is None:
        return np.nan
    if isinstance(v, timedelta):
        return v.total_seconds()
    if hasattr(v, "hour"):
        return v.hour * 3600 + v.minute * 60 + v.second + getattr(v, "microsecond", 0) / 1e6
    return float(v)


_COMMA_NUM = re.compile(r"^[+-]?\d+,\d+$")    # "0,234" / "30,9": decimal-comma text


def _comma_decimal(rows) -> bool:
    """True when the export was written on a decimal-comma locale.

    Such exports store readings as text ("0,234") instead of numbers.
    """
    return any(isinstance(v, str) and _COMMA_NUM.match(v.strip())
               for r in rows for v in r[2:])


def _number(v, comma: bool) -> float | None:
    """A reading as float, or None if it isn't one (blank, OVRFLW, ...).

    In decimal-comma exports Excel sometimes reads "1,236" as one thousand two
    hundred thirty-six (thousands grouping). That needs exactly three digits after a
    non-zero leading digit, so it always lands at >= 1000, far above any real
    absorbance or temperature: undo it.
    """
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v / 1000 if comma and abs(v) >= 1000 else float(v)
    if comma and isinstance(v, str) and _COMMA_NUM.match(v.strip()):
        return float(v.strip().replace(",", "."))
    return None


def _started(meta: dict) -> datetime | None:
    d, t = meta.get("Date"), meta.get("Time")
    if isinstance(t, time):
        day = d.date() if isinstance(d, datetime) else datetime(1900, 1, 1).date()
        return datetime.combine(day, t)
    return None


def read_plate(data: bytes, filename: str) -> Plate:
    """Parse one kinetic export: only the kinetic block and the read's start time are used.

    OVRFLW -> NaN + saturated; empty cell -> NaN + missing. Both numeric exports and
    decimal-comma text exports ("0,234") are accepted; the format is detected per file.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    finally:
        wb.close()

    meta = {r[0]: r[1] for r in rows[:20] if r and isinstance(r[0], str)}
    hdr = next((i for i, r in enumerate(rows)
                if len(r) > 3 and r[1] == "Time" and isinstance(r[2], str)
                and r[2].startswith("T")), None)
    if hdr is None:
        raise ValueError(f"{filename}: no kinetic block (a 'Time' / 'T° …' header) found")

    # Map each well column present in the export onto its slot in WELLS.
    cols = {j: name for j, name in enumerate(rows[hdr]) if j >= 3 and name in WELLS}
    slot = {j: WELLS.index(name) for j, name in cols.items()}

    block = []
    for r in rows[hdr + 1:]:
        if len(r) < 4 or r[1] is None or r[3] is None:
            break                                   # end of the kinetic block
        block.append(r)
    comma = _comma_decimal(block)

    times, temps, vals, sat, miss = [], [], [], [], []
    for r in block:
        times.append(_seconds(r[1]))
        temp = _number(r[2], comma)
        temps.append(np.nan if temp is None else temp)
        v_row, s_row, m_row = np.full(96, np.nan), np.zeros(96, bool), np.ones(96, bool)
        for j, k in slot.items():
            v = r[j] if j < len(r) else None
            x = _number(v, comma)
            if x is not None:
                v_row[k], m_row[k] = x, False
            elif isinstance(v, str) and v.strip().upper().startswith("OVRFLW"):
                s_row[k], m_row[k] = True, False
        vals.append(v_row); sat.append(s_row); miss.append(m_row)

    if not times:
        raise ValueError(f"{filename}: the kinetic block is empty")

    t, temp = np.asarray(times, float), np.asarray(temps, float)
    a, sat, miss = np.vstack(vals), np.vstack(sat), np.vstack(miss)

    # A run stopped mid-sweep leaves a final partial timepoint (some wells blank
    # rather than OVRFLW). Drop trailing timepoints that miss a well the plate has.
    present = np.zeros(96, bool)
    present[list(slot.values())] = True
    dropped = 0
    while len(t) > 1 and miss[-1, present].any():
        t, temp, a, sat, miss = t[:-1], temp[:-1], a[:-1], sat[:-1], miss[:-1]
        dropped += 1

    return Plate(file=filename, t=t, temp=temp, a=a, sat=sat, dropped=dropped,
                 started=_started(meta))
